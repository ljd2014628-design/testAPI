# -*- coding: utf-8 -*-
"""Excel 模板生成、导入解析、结果导出。

Sheet1 接口用例   : 用例清单 + 最近 5 次返回结果
Sheet2 组合场景结果: 组合(多接口)场景最新一次执行结果，按行记录
Sheet3 场景配置   : 组合场景的步骤编排与变量提取
Sheet4 公共配置   : BaseURL / 编码 / 超时 / 公共请求头
"""

import json
import re
import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"

SHEET_CASES = "接口用例"
SHEET_SCENE_RESULT = "组合场景结果"
SHEET_SCENE_CFG = "场景配置"
SHEET_COMMON = "公共配置"

MAX_HISTORY = 5

CASE_HEADERS = [
    "用例ID", "接口名称", "接口用例", "接口URL", "请求方法",
    "接口请求头", "请求Body", "URL参数", "字段备注", "是否执行", "提取变量",
    "返回结果1(最新)", "返回结果2", "返回结果3", "返回结果4", "返回结果5",
    "预期结果内容", "是否符合预期",
]
CASE_WIDTHS = [14, 18, 22, 34, 10, 26, 40, 30, 26, 10, 26, 46, 30, 30, 30, 30, 36, 40]

SCENE_RESULT_HEADERS = [
    "场景名称", "步骤", "用例ID", "接口名称", "请求方法", "完整URL",
    "请求Body", "响应状态", "请求结果", "耗时(ms)", "执行时间", "是否符合预期",
]
SCENE_RESULT_WIDTHS = [20, 6, 14, 18, 10, 46, 40, 10, 60, 10, 20, 36]

SCENE_CFG_HEADERS = ["场景名称", "步骤", "用例ID", "变量提取", "备注"]
SCENE_CFG_WIDTHS = [22, 8, 14, 34, 30]

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
NOTE_FONT = Font(name=FONT, italic=True, color="7F7F7F", size=9)
TITLE_FONT = Font(name=FONT, bold=True, size=12, color="2F5597")
SAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, headers, widths, row=1):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 26
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _cell(ws, r, c, v, sample=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=10)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = BORDER
    if sample:
        cell.fill = SAMPLE_FILL
    return cell


# --------------------------------------------------------------------------
# URL参数 / 字段备注 的文本表示
#   URL参数： 每行 key=value          备注写在  # 后面；行首 # 表示停用该参数
#   字段备注：每行 字段路径 = 备注
# --------------------------------------------------------------------------
def dump_params_text(params):
    lines = []
    for p in params or []:
        if not p.get("key"):
            continue
        line = "%s=%s" % (p["key"], p.get("value", "") or "")
        if p.get("memo"):
            line += "    # " + str(p["memo"]).replace("\n", " ")
        if not p.get("on", True):
            line = "#" + line
        lines.append(line)
    return "\n".join(lines)


def parse_params_text(text):
    params = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        on = True
        if line.startswith("#"):
            on = False
            line = line.lstrip("#").strip()
        memo = ""
        if "#" in line:
            line, memo = line.split("#", 1)
            line, memo = line.strip(), memo.strip()
        if "=" in line:
            k, v = line.split("=", 1)
        else:
            k, v = line, ""
        if k.strip():
            params.append({"key": k.strip(), "value": v.strip(), "memo": memo, "on": on})
    return params


def dump_asserts_text(asserts):
    """每行：目标 [比较方式] 期望值    # 备注；行首 # 表示停用。"""
    lines = []
    for a in asserts or []:
        if not a.get("target"):
            continue
        line = ("%s [%s] %s" % (a["target"], a.get("op", "等于"),
                                a.get("expected", "") or "")).rstrip()
        if a.get("level") == "仅警告":
            line += " @警告"
        if a.get("memo"):
            line += "    # " + str(a["memo"]).replace("\n", " ")
        if not a.get("on", True):
            line = "#" + line
        lines.append(line)
    return "\n".join(lines)


def parse_asserts_text(text):
    out = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        on = True
        if line.startswith("#"):
            on = False
            line = line.lstrip("#").strip()
        m = re.match(r"^(.*?)\s*\[(.+?)\]\s*(.*)$", line)
        if not m:
            continue
        target, op, rest = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
        memo = ""
        if "#" in rest:
            rest, memo = rest.split("#", 1)
            rest, memo = rest.strip(), memo.strip()
        level = "必须通过"
        if "@警告" in rest:
            level = "仅警告"
            rest = rest.replace("@警告", "").strip()
        if target:
            out.append({"on": on, "target": target, "op": op, "expected": rest,
                        "level": level, "memo": memo})
    return out


def assert_summary(results):
    """最近一次执行的断言结论，写进「是否符合预期」列。"""
    if not results:
        return ""
    res = results[0]
    if res.get("imported"):
        return ""
    rows = res.get("asserts") or []
    if not rows:
        return "未设置断言"
    fail = [a for a in rows if not a.get("ok") and a.get("level") != "仅警告"]
    warn = [a for a in rows if not a.get("ok") and a.get("level") == "仅警告"]
    if fail:
        head = "不符合预期（%d 条未通过" % len(fail) + ("，%d 条警告）" % len(warn) if warn else "）")
    elif warn:
        head = "符合预期（%d 条通过，%d 条仅警告）" % (len(rows) - len(warn), len(warn))
    else:
        head = "符合预期（%d 条全部通过）" % len(rows)
    lines = [head]
    for a in rows:
        mark = "√" if a.get("ok") else ("⚠" if a.get("level") == "仅警告" else "×")
        lines.append("%s %s [%s] %s → %s" % (
            mark, a.get("target", ""), a.get("op", ""), a.get("expected", ""), a.get("msg", "")))
    return "\n".join(lines)


def dump_memos_text(memos):
    return "\n".join("%s = %s" % (k, str(v).replace("\n", " "))
                     for k, v in (memos or {}).items() if v)


def parse_memos_text(text):
    out = {}
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line or "=" not in line:
            continue
        k, v = line.split("=", 1)
        if k.strip() and v.strip():
            out[k.strip()] = v.strip()
    return out


# --------------------------------------------------------------------------
# 模板
# --------------------------------------------------------------------------
SAMPLE_CASES = [
    {
        "id": "CASE_001",
        "name": "查询可用营销活动",
        "caseName": "正常下单-带优惠券",
        "url": "/marketing/promotion/findAll",
        "method": "POST",
        "headers": "",
        "body": json.dumps({
            "orderNo": "LS20260912000005",
            "customerPhone": "15715852024",
            "storeCode": "80001224",
            "companyId": 1,
            "brandCode": "K001",
            "scene": 1,
            "lines": [{"lineId": "30", "outerSkuId": "90.8101.C111108",
                       "qty": 1, "unitPrice": 230000}],
            "selectedResources": {"couponCodes": ["03321876032185475072"]},
        }, ensure_ascii=False, indent=2),
        "params": "",
        "memos": "orderNo = 订单号，每次要换\nlines[0].unitPrice = 单价，单位：分",
        "asserts": "状态码 [等于] 200    # HTTP 要正常\n$.code [等于] 0    # 业务码正常\n$.data.promotionId [不为空]    # 必须返回活动ID",
        "enabled": "Y",
        "extract": "couponInstanceId=$.data.lines[0].couponInstanceId\npromotionId=$.data.promotionId",
    },
    {
        "id": "CASE_002",
        "name": "查询优惠券详情",
        "caseName": "用上一步返回的券ID查询",
        "url": "/coupon/query/${couponInstanceId}",
        "method": "GET",
        "headers": "X-Trace-Id: local-debug-001",
        "body": "",
        "asserts": "$.data.remain [较上次减少] 1    # 领券场景：每跑一次少一张，不用改期望值",
        "params": "brandCode=K001    # 品牌编码，不填查不到\npageNo=1    # 第几页\n#debug=true    # 行首加#表示这个参数本次不发送",
        "memos": "",
        "enabled": "Y",
        "extract": "couponStatus=$.data.status",
    },
    {
        "id": "CASE_003",
        "name": "营销活动变更",
        "caseName": "变更已选活动",
        "url": "/marketing/promotion/change",
        "method": "POST",
        "headers": json.dumps({"Content-Type": "application/json",
                               "tenant": "K001"}, ensure_ascii=False),
        "body": json.dumps({
            "orderNo": "LS20260912000005",
            "promotionId": "${promotionId}",
            "couponInstanceId": "${couponInstanceId}",
        }, ensure_ascii=False, indent=2),
        "params": "",
        "memos": "",
        "asserts": "",
        "enabled": "N",
        "extract": "",
    },
]

SAMPLE_SCENES = [
    ("下单并核验优惠券", 1, "CASE_001", "couponInstanceId=$.data.lines[0].couponInstanceId", "取第一步返回的券实例ID"),
    ("下单并核验优惠券", 2, "CASE_002", "couponStatus=$.data.status", "URL 中用 ${couponInstanceId} 引用上一步的值"),
]


def build_template(path):
    wb = Workbook()

    # ---------------- Sheet1 接口用例 ----------------
    ws = wb.active
    ws.title = SHEET_CASES
    _style_header(ws, CASE_HEADERS, CASE_WIDTHS)
    r = 2
    for s in SAMPLE_CASES:
        _cell(ws, r, 1, s["id"], True)
        _cell(ws, r, 2, s["name"], True)
        _cell(ws, r, 3, s["caseName"], True)
        _cell(ws, r, 4, s["url"], True)
        _cell(ws, r, 5, s["method"], True)
        _cell(ws, r, 6, s["headers"], True)
        _cell(ws, r, 7, s["body"], True)
        _cell(ws, r, 8, s.get("params", ""), True)
        _cell(ws, r, 9, s.get("memos", ""), True)
        _cell(ws, r, 10, s["enabled"], True)
        _cell(ws, r, 11, s["extract"], True)
        for c in range(12, 17):
            _cell(ws, r, c, None)
        _cell(ws, r, 17, s.get("asserts", ""), True)
        _cell(ws, r, 18, None)
        ws.row_dimensions[r].height = 90
        r += 1
    for extra in range(r, r + 20):           # 预留空行
        for c in range(1, 19):
            _cell(ws, extra, c, None)

    dv = DataValidation(type="list", formula1='"GET,POST,PUT,DELETE,PATCH"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("E2:E%d" % (r + 20))
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("J2:J%d" % (r + 20))

    note_row = r + 22
    ws.cell(row=note_row, column=1, value="填写说明（黄色底为示例行，可直接改写或删除）").font = TITLE_FONT
    notes = [
        "用例ID：唯一标识，组合场景通过它引用该接口，建议 CASE_001 递增，不要重复。",
        "接口URL：可填相对路径（自动拼接公共配置里的 BaseURL），也可填 http 开头的完整地址。",
        "接口请求头：只填该接口特殊的头。支持两种写法 —— ① JSON：{\"token\":\"abc\"} ② 每行一个：token: abc。",
        "            读取优先级：接口请求头 > 公共配置请求头（同名时接口级覆盖公共）。",
        "请求Body：POST/PUT 填 JSON；GET 的查询参数建议填到「URL参数」列，不用拼在 URL 上。",
        "URL参数：每行一个 key=value，会自动拼到 URL 后面；# 后面写备注；行首加 # 表示这行本次不发送。",
        "        例：brandCode=K001    # 品牌编码    值里同样可以用 ${变量名}。",
        "字段备注：给 Body 里的字段加注释，每行 字段路径 = 备注，例：lines[0].unitPrice = 单价，单位分。",
        "预期结果内容（断言）：每行一条  目标 [比较方式] 期望值  # 备注；行首加 # 表示这条先不校验。",
        "        目标可以是 JSON 路径（$.data.remain）、状态码、耗时ms、响应文本。",
        "        比较方式：等于/不等于/大于/小于/包含/正则匹配/存在/不存在/为空/不为空/长度等于，",
        "        以及「较上次增加|较上次减少|较上次不变|较上次变化」——用于领券这类每跑一次就变的场景。",
        "        例：$.data.remain [较上次减少] 1  → 10→9 通过，9→8 也通过，期望值不用每次改。",
        "是否符合预期：工具执行后自动回填，无需手工填写。",
        "        备注只是给自己看的，发请求时不会带上。",
        "是否执行：Y 参与「全部执行」，N 跳过（在工具里仍可单独勾选执行）。",
        "提取变量：把本接口的响应值存成变量，供后续接口引用。每行一个，格式 变量名=$.json.路径",
        "          例：couponId=$.data.lines[0].couponInstanceId ；数组下标用 [0]。",
        "引用变量：在 URL / Body / 请求头里写 ${变量名} 即可，例 /coupon/query/${couponId}。",
        "返回结果1-5：工具执行后自动回填，返回结果1 永远是最近一次，最多保留 5 次，无需手工填写。",
    ]
    for i, t in enumerate(notes):
        c = ws.cell(row=note_row + 1 + i, column=1, value=t)
        c.font = NOTE_FONT

    # ---------------- Sheet2 组合场景结果 ----------------
    ws2 = wb.create_sheet(SHEET_SCENE_RESULT)
    _style_header(ws2, SCENE_RESULT_HEADERS, SCENE_RESULT_WIDTHS)
    ws2.cell(row=3, column=1,
             value="本页由工具自动写入：多接口组合场景的执行结果，按行记录，只保留最新一次执行。").font = NOTE_FONT

    # ---------------- Sheet3 场景配置 ----------------
    ws3 = wb.create_sheet(SHEET_SCENE_CFG)
    _style_header(ws3, SCENE_CFG_HEADERS, SCENE_CFG_WIDTHS)
    r = 2
    for name, step, cid, ext, memo in SAMPLE_SCENES:
        _cell(ws3, r, 1, name, True)
        _cell(ws3, r, 2, step, True)
        _cell(ws3, r, 3, cid, True)
        _cell(ws3, r, 4, ext, True)
        _cell(ws3, r, 5, memo, True)
        ws3.row_dimensions[r].height = 30
        r += 1
    for extra in range(r, r + 20):
        for c in range(1, 6):
            _cell(ws3, extra, c, None)
    nr = r + 2
    ws3.cell(row=nr, column=1, value="填写说明").font = TITLE_FONT
    for i, t in enumerate([
        "同一个「场景名称」的多行 = 一个组合场景，按「步骤」从小到大依次执行。",
        "用例ID：引用 Sheet1《接口用例》里的用例，必须存在。",
        "变量提取：本步骤执行后从响应里取值，格式 变量名=$.路径，每行一个；后续步骤用 ${变量名} 引用。",
        "         留空则沿用 Sheet1 该用例自带的「提取变量」配置。",
        "场景也可以直接在工具界面上新建、编辑，保存后导出会写回本页。",
    ]):
        ws3.cell(row=nr + 1 + i, column=1, value=t).font = NOTE_FONT

    # ---------------- Sheet4 公共配置 ----------------
    ws4 = wb.create_sheet(SHEET_COMMON)
    ws4.column_dimensions["A"].width = 22
    ws4.column_dimensions["B"].width = 46
    ws4.column_dimensions["C"].width = 12
    ws4.cell(row=1, column=1, value="基础配置").font = TITLE_FONT
    base = [("BaseURL", "https://api-test.example.com", "接口URL为相对路径时自动拼接"),
            ("编码", "utf-8", "默认 utf-8"),
            ("超时(秒)", 30, "单个请求超时时间"),
            ("校验SSL", "N", "Y=校验证书，N=忽略（测试环境自签证书填N）")]
    r = 2
    for k, v, memo in base:
        ws4.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        c = ws4.cell(row=r, column=2, value=v)
        c.font = Font(name=FONT, size=10)
        c.fill = SAMPLE_FILL
        ws4.cell(row=r, column=3, value=memo).font = NOTE_FONT
        r += 1

    r += 1
    ws4.cell(row=r, column=1, value="公共请求头（所有接口生效，可被接口级请求头覆盖）").font = TITLE_FONT
    r += 1
    for i, h in enumerate(["头名称", "头值", "启用(Y/N)"]):
        c = ws4.cell(row=r, column=i + 1, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center")
    r += 1
    for k, v in [("Content-Type", "application/json;charset=utf-8"),
                 ("Accept", "application/json"),
                 ("Authorization", "Bearer PUT-YOUR-TOKEN-HERE"),
                 ("brandCode", "K001")]:
        _cell(ws4, r, 1, k, True)
        _cell(ws4, r, 2, v, True)
        _cell(ws4, r, 3, "Y", True)
        r += 1
    for extra in range(r, r + 10):
        for c in range(1, 4):
            _cell(ws4, extra, c, None)

    wb.save(path)
    return path


# --------------------------------------------------------------------------
# 导入
# --------------------------------------------------------------------------
def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip() if isinstance(v, str) else str(v)


def _row_empty(ws, row, ncol):
    return all(_s(ws.cell(row=row, column=c).value) == "" for c in range(1, ncol + 1))


def read_workbook(path):
    """解析工作簿，返回 dict(config, cases, scenarios)。兼容旧版/简版表头。"""
    wb = load_workbook(path, data_only=True)
    result = {"config": None, "cases": [], "scenarios": []}

    # ---- 公共配置 ----
    if SHEET_COMMON in wb.sheetnames:
        ws = wb[SHEET_COMMON]
        cfg = {"baseUrl": "", "encoding": "utf-8", "timeout": 30,
               "verifySsl": False, "headers": []}
        in_headers = False
        for row in range(1, ws.max_row + 1):
            a = _s(ws.cell(row=row, column=1).value)
            b = _s(ws.cell(row=row, column=2).value)
            c = _s(ws.cell(row=row, column=3).value)
            if not a:
                continue
            if a == "BaseURL":
                cfg["baseUrl"] = b
            elif a == "编码":
                cfg["encoding"] = b or "utf-8"
            elif a == "超时(秒)":
                try:
                    cfg["timeout"] = int(float(b))
                except Exception:
                    cfg["timeout"] = 30
            elif a == "校验SSL":
                cfg["verifySsl"] = b.upper().startswith("Y")
            elif a == "头名称":
                in_headers = True
            elif in_headers and a not in ("公共请求头", ""):
                if a.startswith("公共请求头"):
                    continue
                cfg["headers"].append({"key": a, "value": b,
                                       "enabled": (c.upper() != "N")})
        result["config"] = cfg

    # ---- 接口用例 ----
    ws = None
    for nm in (SHEET_CASES, "Sheet1"):
        if nm in wb.sheetnames:
            ws = wb[nm]
            break
    if ws is None:
        ws = wb.worksheets[0]

    headers = [_s(ws.cell(row=1, column=c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    def col(*names):
        for n in names:
            if n in idx:
                return idx[n]
        return None

    c_id = col("用例ID")
    c_name = col("接口名称")
    c_case = col("接口用例", "用例名称")
    c_url = col("接口URL", "URL")
    c_method = col("接口请求方法", "请求方法")
    c_head = col("接口请求头", "请求头")
    c_body = col("请求Body", "body", "Body", "请求参数")
    c_en = col("是否执行")
    c_ext = col("提取变量")
    c_par = col("URL参数")
    c_memo = col("字段备注")
    c_assert = col("预期结果内容")
    res_cols = []
    for i in range(1, MAX_HISTORY + 1):
        cc = col("返回结果%d(最新)" % i, "返回结果%d" % i)
        if cc:
            res_cols.append(cc)

    seq = 0
    for row in range(2, ws.max_row + 1):
        if _row_empty(ws, row, ws.max_column):
            continue
        url = _s(ws.cell(row=row, column=c_url).value) if c_url else ""
        name = _s(ws.cell(row=row, column=c_name).value) if c_name else ""
        if not url and not name:
            continue
        if name == "填写说明（黄色底为示例行，可直接改写或删除）" or (not url and not (c_body and _s(ws.cell(row=row, column=c_body).value))):
            # 说明区
            if not url:
                continue
        seq += 1
        cid = _s(ws.cell(row=row, column=c_id).value) if c_id else ""
        if not cid:
            cid = "CASE_%03d" % seq
        case = {
            "id": cid,
            "name": name or cid,
            "caseName": _s(ws.cell(row=row, column=c_case).value) if c_case else "",
            "url": url,
            "method": (_s(ws.cell(row=row, column=c_method).value) or "GET").upper() if c_method else "GET",
            "headers": _s(ws.cell(row=row, column=c_head).value) if c_head else "",
            "body": _s(ws.cell(row=row, column=c_body).value) if c_body else "",
            "enabled": (_s(ws.cell(row=row, column=c_en).value).upper() != "N") if c_en else True,
            "extract": _s(ws.cell(row=row, column=c_ext).value) if c_ext else "",
            "params": parse_params_text(_s(ws.cell(row=row, column=c_par).value)) if c_par else [],
            "memos": parse_memos_text(_s(ws.cell(row=row, column=c_memo).value)) if c_memo else {},
            "asserts": parse_asserts_text(_s(ws.cell(row=row, column=c_assert).value)) if c_assert else [],
            "results": [],
        }
        for cc in res_cols:
            txt = _s(ws.cell(row=row, column=cc).value)
            if txt:
                case["results"].append({"text": txt, "imported": True,
                                        "time": "", "ok": None, "status": ""})
        result["cases"].append(case)

    # ---- 场景配置 ----
    if SHEET_SCENE_CFG in wb.sheetnames:
        ws3 = wb[SHEET_SCENE_CFG]
        buckets = {}
        order = []
        for row in range(2, ws3.max_row + 1):
            nm = _s(ws3.cell(row=row, column=1).value)
            cid = _s(ws3.cell(row=row, column=3).value)
            if not nm or not cid:
                continue
            if nm == "填写说明":
                break
            try:
                step = int(float(_s(ws3.cell(row=row, column=2).value) or 0))
            except Exception:
                step = 0
            if nm not in buckets:
                buckets[nm] = []
                order.append(nm)
            buckets[nm].append({
                "step": step,
                "caseId": cid,
                "extract": _s(ws3.cell(row=row, column=4).value),
                "memo": _s(ws3.cell(row=row, column=5).value),
            })
        for i, nm in enumerate(order):
            steps = sorted(buckets[nm], key=lambda x: x["step"])
            result["scenarios"].append({
                "id": "SC_%03d" % (i + 1),
                "name": nm,
                "enabled": True,
                "steps": [{"caseId": s["caseId"], "extract": s["extract"], "memo": s["memo"]}
                          for s in steps],
                "results": [],
            })
    return result


# --------------------------------------------------------------------------
# 导出
# --------------------------------------------------------------------------
def write_workbook(path, state):
    cfg = state["config"]
    wb = Workbook()

    ws = wb.active
    ws.title = SHEET_CASES
    _style_header(ws, CASE_HEADERS, CASE_WIDTHS)
    r = 2
    for case in state["cases"]:
        _cell(ws, r, 1, case["id"])
        _cell(ws, r, 2, case["name"])
        _cell(ws, r, 3, case.get("caseName", ""))
        _cell(ws, r, 4, case["url"])
        _cell(ws, r, 5, case["method"])
        _cell(ws, r, 6, case.get("headers", ""))
        _cell(ws, r, 7, case.get("body", ""))
        _cell(ws, r, 8, dump_params_text(case.get("params")))
        _cell(ws, r, 9, dump_memos_text(case.get("memos")))
        _cell(ws, r, 10, "Y" if case.get("enabled", True) else "N")
        _cell(ws, r, 11, case.get("extract", ""))
        for i in range(MAX_HISTORY):
            txt = ""
            if i < len(case.get("results", [])):
                txt = result_to_text(case["results"][i])
            _cell(ws, r, 12 + i, txt)
        _cell(ws, r, 17, dump_asserts_text(case.get("asserts")))
        _cell(ws, r, 18, assert_summary(case.get("results")))
        ws.row_dimensions[r].height = 80
        r += 1

    # Sheet2 组合场景结果（只保留最新一次，按行记录）
    ws2 = wb.create_sheet(SHEET_SCENE_RESULT)
    _style_header(ws2, SCENE_RESULT_HEADERS, SCENE_RESULT_WIDTHS)
    r = 2
    for sc in state["scenarios"]:
        results = sc.get("results", [])
        if not results:
            continue
        latest = results[0]
        for i, st in enumerate(latest.get("steps", []), start=1):
            _cell(ws2, r, 1, sc["name"])
            _cell(ws2, r, 2, i)
            _cell(ws2, r, 3, st.get("caseId", ""))
            _cell(ws2, r, 4, st.get("caseName", ""))
            _cell(ws2, r, 5, st.get("method", ""))
            _cell(ws2, r, 6, st.get("url", ""))
            _cell(ws2, r, 7, st.get("reqBody", ""))
            _cell(ws2, r, 8, st.get("status", ""))
            _cell(ws2, r, 9, st.get("respBody", "") or st.get("error", ""))
            _cell(ws2, r, 10, st.get("elapsed", ""))
            _cell(ws2, r, 11, latest.get("time", ""))
            _cell(ws2, r, 12, assert_summary([st]))
            ws2.row_dimensions[r].height = 60
            r += 1

    # Sheet3 场景配置
    ws3 = wb.create_sheet(SHEET_SCENE_CFG)
    _style_header(ws3, SCENE_CFG_HEADERS, SCENE_CFG_WIDTHS)
    r = 2
    for sc in state["scenarios"]:
        for i, st in enumerate(sc.get("steps", []), start=1):
            _cell(ws3, r, 1, sc["name"])
            _cell(ws3, r, 2, i)
            _cell(ws3, r, 3, st.get("caseId", ""))
            _cell(ws3, r, 4, st.get("extract", ""))
            _cell(ws3, r, 5, st.get("memo", ""))
            r += 1

    # Sheet4 公共配置
    ws4 = wb.create_sheet(SHEET_COMMON)
    ws4.column_dimensions["A"].width = 22
    ws4.column_dimensions["B"].width = 46
    ws4.column_dimensions["C"].width = 12
    ws4.cell(row=1, column=1, value="基础配置").font = TITLE_FONT
    rows = [("BaseURL", cfg.get("baseUrl", "")),
            ("编码", cfg.get("encoding", "utf-8")),
            ("超时(秒)", cfg.get("timeout", 30)),
            ("校验SSL", "Y" if cfg.get("verifySsl") else "N")]
    r = 2
    for k, v in rows:
        ws4.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        ws4.cell(row=r, column=2, value=v).font = Font(name=FONT, size=10)
        r += 1
    r += 1
    ws4.cell(row=r, column=1, value="公共请求头（所有接口生效，可被接口级请求头覆盖）").font = TITLE_FONT
    r += 1
    for i, h in enumerate(["头名称", "头值", "启用(Y/N)"]):
        c = ws4.cell(row=r, column=i + 1, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
    r += 1
    for h in cfg.get("headers", []):
        _cell(ws4, r, 1, h.get("key", ""))
        _cell(ws4, r, 2, h.get("value", ""))
        _cell(ws4, r, 3, "Y" if h.get("enabled", True) else "N")
        r += 1

    wb.save(path)
    return path


def result_to_text(res):
    txt = _result_to_text(res)
    if len(txt) > 30000:  # Excel 单元格上限 32767
        txt = txt[:30000] + "\n...(内容过长，Excel 单元格已截断；完整内容请在工具里查看)"
    return txt


def _result_to_text(res):
    if res.get("imported") and res.get("text"):
        return res["text"]
    lines = [
        "【%s】%s  状态码:%s  耗时:%sms" % (
            res.get("time", ""),
            "成功" if res.get("ok") else "失败",
            res.get("status", "-"), res.get("elapsed", "-")),
        "请求URL: %s %s" % (res.get("method", ""), res.get("url", "")),
    ]
    if res.get("reqHeaders"):
        lines.append("请求头: " + res["reqHeaders"])
    if res.get("reqBody"):
        lines.append("请求Body: " + res["reqBody"])
    if res.get("error"):
        lines.append("异常: " + res["error"])
    else:
        lines.append("响应: " + (res.get("respBody") or ""))
    if res.get("vars"):
        lines.append("提取变量: " + res["vars"])
    return "\n".join(lines)


def now_str():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
